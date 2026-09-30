"""Unit tests for the sample generator's files, fixtures and manifest; needs no services."""

import json
from datetime import datetime, timezone

import pytest
from docx import Document
from openpyxl import load_workbook

from generate_samples import file_uri, generate, parse_offset

NOW = datetime(2026, 9, 29, 12, 0, tzinfo=timezone.utc)
BASE = "http://sharepoint-mock:8000"


@pytest.fixture(scope="module")
def generated(tmp_path_factory):
    """Generate the samples into a temp directory and return the root with the manifest."""
    root = tmp_path_factory.mktemp("samples")
    manifest = generate(root / "files", root / "fixtures", BASE, now=NOW)
    return root, manifest


def docx_lines(path):
    """Return each paragraph of a docx as a (style name, text) pair."""
    return [(p.style.name, p.text) for p in Document(path).paragraphs]


def sample(root, server_path):
    """Return the generated file path for a server-relative path."""
    return root / "files" / server_path.lstrip("/")


def test_every_manifest_file_exists_and_opens(generated):
    """Every manifest file exists and opens as docx or xlsx."""
    root, manifest = generated
    for entry in manifest["files"]:
        path = sample(root, entry["path"])
        assert path.exists(), entry["path"]
        if path.suffix == ".docx":
            Document(path)
        else:
            load_workbook(path, read_only=True)


def test_deliberate_absences(generated):
    """The missing runbook, the 3.0.0 approval and the 3.0.0 gate fixture are not generated."""
    root, _ = generated
    assert not sample(root, "/sites/compliance/Shared Documents/Runbooks/missing.docx").exists()
    assert not sample(root, "/sites/cds/Approvals/AIT-12345-3.0.0-approval.docx").exists()
    assert not (root / "fixtures/ci/deploy-AIT-12345-3.0.0.json").exists()


def test_freshness_metadata_offsets(generated):
    """The runbooks are 3 and 90 days old, and the fresh one has its owner team."""
    _, manifest = generated
    by_path = {e["path"]: e for e in manifest["files"]}
    fresh = by_path["/sites/compliance/Shared Documents/Runbooks/runbook.docx"]
    stale = by_path["/sites/compliance/Shared Documents/Runbooks/stale-runbook.docx"]
    assert (NOW - parse_offset(fresh["modified"], NOW)).days == 3
    assert (NOW - parse_offset(stale["modified"], NOW)).days == 90
    assert fresh["list_fields"]["Owner_x0020_Team"] == "Platform Operations"


def test_cloud_register_rows_and_as_of(generated):
    """The September register lists AIT-12345 Yes and AIT-67890 No with its as-of date."""
    root, manifest = generated
    by_path = {e["path"]: e for e in manifest["files"]}
    path = "/sites/compliance/Shared Documents/APS/cloud-approvals-2026-09.xlsx"
    assert by_path[path]["modified"] == "-5d"
    wb = load_workbook(sample(root, path), read_only=True)
    rows = list(wb["AITs"].iter_rows(values_only=True))
    assert rows[0] == ("Application ID", "Application Name", "Approved for Cloud")
    approvals = {r[0]: r[2] for r in rows[1:]}
    assert approvals == {"AIT-12345": "Yes", "AIT-67890": "No"}
    assert "AIT-99999" not in approvals
    assert wb["Info"]["B2"].value == "2026-09-24"


def test_older_cloud_register_disagrees(generated):
    """The June register marks AIT-12345 as not approved."""
    root, _ = generated
    wb = load_workbook(sample(root, "/sites/compliance/Shared Documents/APS/cloud-approvals-2026-06.xlsx"), read_only=True)
    approvals = {r[0]: r[2] for r in list(wb["AITs"].iter_rows(values_only=True))[1:]}
    assert approvals["AIT-12345"] == "No"


@pytest.mark.parametrize("version,approver,signed", [
    ("2.4.0", "A. Approver", "2026-09-22"),
    ("2.5.0", "Z. Outsider", "2026-09-24"),
    ("2.6.0", "B. Approver", "2026-09-27"),
])
def test_cds_approvals(generated, version, approver, signed):
    """Each approval doc has status, approver and sign-off, and only 2.5.0's approver is not an owner."""
    root, manifest = generated
    lines = docx_lines(sample(root, f"/sites/cds/Approvals/AIT-12345-{version}-approval.docx"))
    texts = [t for _, t in lines]
    assert ("Heading 2", "Approval Status") in lines
    assert "Approved for production deployment" in texts
    assert f"Approved by: {approver}" in texts
    assert f"Sign-off date: {signed}" in texts
    members = {manifest["users"][m]["Title"] for m in manifest["sites"]["/sites/cds"]["owner_group"]["members"]}
    assert (approver in members) is (version != "2.5.0")


@pytest.mark.parametrize("exc,expiry,has_controls", [
    ("EXC-001", "2026-11-28", True),
    ("EXC-002", "2026-09-19", True),
    ("EXC-003", "2026-11-28", False),
])
def test_exceptions(generated, exc, expiry, has_controls):
    """Each exception has its expiry, and only EXC-003 lacks a Compensating Controls section."""
    root, _ = generated
    lines = docx_lines(sample(root, f"/sites/compliance/Shared Documents/Exceptions/{exc}.docx"))
    assert f"Expiry date: {expiry}" in [t for _, t in lines]
    assert (("Heading 2", "Compensating Controls") in lines) is has_controls


def test_gate_fixtures_reference_approvals(generated):
    """Gate fixtures allow, and only 2.6.0 points to a different version's approval."""
    root, _ = generated
    def gate(version):
        """Return the gate object from a version's deploy fixture."""
        return json.loads((root / f"fixtures/ci/deploy-AIT-12345-{version}.json").read_text())["gate"]

    for version in ("2.4.0", "2.5.0"):
        assert gate(version)["approval_uri"] == file_uri(BASE, f"/sites/cds/Approvals/AIT-12345-{version}-approval.docx")
        assert gate(version)["result"] == "allow"
    assert gate("2.6.0")["approval_uri"] != file_uri(BASE, "/sites/cds/Approvals/AIT-12345-2.6.0-approval.docx")


def test_uri_encodes_spaces():
    """file_uri percent-encodes spaces in the path."""
    assert file_uri(BASE, "/sites/compliance/Shared Documents/a.docx") == \
        "http://sharepoint-mock:8000/sites/compliance/Shared%20Documents/a.docx"


def test_bad_offset_rejected():
    """An offset not in the -Nd form raises ValueError."""
    with pytest.raises(ValueError):
        parse_offset("3 days", NOW)
