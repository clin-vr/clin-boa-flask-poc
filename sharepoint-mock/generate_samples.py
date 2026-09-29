"""Build the mock's sample files and CI fixtures from seed/manifest.json."""

from __future__ import annotations

import argparse
import json
import os
import re
from datetime import datetime, timedelta, timezone
from pathlib import Path
from typing import Any
from urllib.parse import quote

from docx import Document
from openpyxl import Workbook

MANIFEST = Path(__file__).parent / "seed" / "manifest.json"
TOKEN = re.compile(r"\{(date|ts|uri):([^}]+)\}")
OFFSET = re.compile(r"^([+-]\d+)d$")


def parse_offset(offset: str, now: datetime) -> datetime:
    match = OFFSET.match(offset)
    if not match:
        raise ValueError(f"Bad offset {offset!r}; expected like '-3d' or '+60d'")
    return now + timedelta(days=int(match.group(1)))


def file_uri(base_url: str, server_relative_path: str) -> str:
    return base_url.rstrip("/") + quote(server_relative_path, safe="/")


def resolve_tokens(value: Any, now: datetime, base_url: str) -> Any:
    if isinstance(value, dict):
        return {k: resolve_tokens(v, now, base_url) for k, v in value.items()}
    if isinstance(value, list):
        return [resolve_tokens(v, now, base_url) for v in value]
    if not isinstance(value, str):
        return value

    def replace(match: re.Match) -> str:
        kind, arg = match.groups()
        if kind == "uri":
            return file_uri(base_url, arg)
        moment = parse_offset(arg, now)
        return moment.date().isoformat() if kind == "date" else moment.strftime("%Y-%m-%dT%H:%M:%SZ")

    return TOKEN.sub(replace, value)


def build_docx(blocks: list[dict]) -> Document:
    document = Document()
    for block in blocks:
        if "heading" in block:
            document.add_heading(block["heading"], level=block["level"])
        else:
            document.add_paragraph(block["text"])
    return document


def build_xlsx(sheets: dict[str, list[list[Any]]]) -> Workbook:
    workbook = Workbook()
    workbook.remove(workbook.active)
    for name, rows in sheets.items():
        sheet = workbook.create_sheet(name)
        for row in rows:
            sheet.append(row)
    return workbook


def load_manifest(path: Path = MANIFEST) -> dict:
    return json.loads(path.read_text(encoding="utf-8"))


def generate(out_dir: Path, fixtures_dir: Path, base_url: str, now: datetime | None = None,
             manifest: dict | None = None) -> dict:
    now = now or datetime.now(timezone.utc)
    manifest = manifest or load_manifest()

    for entry in manifest["files"]:
        content = resolve_tokens(entry["content"], now, base_url)
        target = out_dir / entry["path"].lstrip("/")
        target.parent.mkdir(parents=True, exist_ok=True)
        if content["type"] == "docx":
            build_docx(content["blocks"]).save(target)
        elif content["type"] == "xlsx":
            build_xlsx(content["sheets"]).save(target)
        else:
            raise ValueError(f"Unsupported sample type {content['type']!r} for {entry['path']}")

    for fixture in manifest["fixtures"]:
        target = fixtures_dir / fixture["path"]
        target.parent.mkdir(parents=True, exist_ok=True)
        body = resolve_tokens(fixture["body"], now, base_url)
        target.write_text(json.dumps(body, indent=2) + "\n", encoding="utf-8")

    return manifest


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--out", type=Path, default=Path(os.environ.get("SAMPLES_DIR", "/data/files")))
    parser.add_argument("--fixtures", type=Path, default=Path(os.environ.get("FIXTURES_DIR", "/fixtures")))
    parser.add_argument("--base-url", default=os.environ.get("SHAREPOINT_BASE_URL", "http://sharepoint-mock:8000"))
    args = parser.parse_args()
    manifest = generate(args.out, args.fixtures, args.base_url)
    print(f"Generated {len(manifest['files'])} files and {len(manifest['fixtures'])} fixtures")


if __name__ == "__main__":
    main()
